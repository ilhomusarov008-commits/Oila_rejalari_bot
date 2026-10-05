import os
import sqlite3
import ast
import re
import math
from datetime import datetime
from aiohttp import web
from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command, CommandStart
from aiogram.types import Message, CallbackQuery, FSInputFile
from aiogram.utils.keyboard import InlineKeyboardBuilder
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.context import FSMContext
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. BOT SOZLAMALARI
TOKEN = os.getenv("BOT_TOKEN")
ADMIN_ID = int(os.getenv("ADMIN_ID", "8671467518"))
DB_PATH = os.getenv("DB_PATH", "oila_va_nikoh_tadqiqoti.db")
EXPORT_PATH = os.getenv("EXPORT_PATH", "Oila_va_nikoh_statistik_tahlil.xlsx")
if not TOKEN:
    raise RuntimeError("BOT_TOKEN environment variable o'rnatilmagan!")

# 2. DEMOGRAFIK SAVOLLAR
GENERAL_QUESTIONS = [("age_group", "Yoshingiz qaysi oraliqda?", ["17–19", "20–22", "23–25",
    "26 va undan katta"]),
    ("course", "Kursingiz?", ["1-kurs", "2-kurs", "3-kurs", "4-kurs", "Magistratura / Boshqa"]),
    ("residence", "Doimiy yashash joyingiz?", ["Shahar", "Tuman/Qishloq"]),
    ("marital_status", "Hozirgi oilaviy holatingiz?", ["Turmush qurmagan", "Turmush qurgan"]),
    ("desired_children", "Kelajakda nechta farzandli bo‘lishni xohlaysiz?", ["0", "1", "2", "3",qkoh qurishning " "maqbul yoshi nechada?\n\n"
            "Javobingizni faqat raqam bilan yozing.\n" "Masalan: 25")
        return
    key, question, options = GENERAL_QUESTIONS[index]
    await state.update_data(general_index=index)
    await state.set_state(Survey.general)
    await message.answer(f"DEMOGRAFIK SAVOL " f"{index + 1}/{len(GENERAL_QUESTIONS)}\n\n"
        f"{question}", reply_markup=option_buttons(options, "general"))

async def send_likert(message, state, index):
    if index >= len(CRITERIA):
        await state.update_data(top5_selected=[])
        await state.set_state(Survey.top5)
        await message.answer("TOP-5 TANLOV\n\n" "Yuqoridagi 25 ta test mezonidan "
            "siz uchun eng muhim 5 tasini tanlang.\n\n" "Birinchi bosganingiz — 1-o‘rin,\n"
            "ikkinchi bosganingiz — 2-o‘rin,\n" "shu tartibda 5-o‘ringacha.\n\n"
            "Tanlangan: 0/5", reply_markup=top5_keyboard([]))
        return
    data = await state.get_data()
    target = target_label(data.get("gender"))
    key, label = CRITERIA[index]
    await state.update_data(likert_index=index)
    await state.set_state(Survey.likert)
    await message.answer(f"TEST SAVOLI " f"{index + 1}/25\n\n" f"Bo‘lajak {target}ning "
        "quyidagi xususiyati siz uchun " "qanchalik muhim?\n\n" f"{label}", reply_markup=likert_buttons(
        ))

async def send_sd(message, state, index):
    if index >= len(SD_QUESTIONS):
        await state.set_state(Survey.main_factor)
        await message.answer("YAKUNIY SAVOL\n\n" "Umuman olganda, juft tanlashda "
            "siz uchun qaysi umumiy yo‘nalish " "eng muhim?", reply_markup=option_buttons(
            MAIN_FACTORS, "factor"))
        return
    key, question, expected = SD_QUESTIONS[index]
    await state.update_data(sd_index=index)
    await state.set_state(Survey.sd_scale)
    await message.answer(f"SD SAVOLI " f"{index + 1}/{len(SD_QUESTIONS)}\n\n" f"{question}",
        reply_markup=sd_buttons())

# 11. EXCEL EKSPORT

def style_sheet(ws):
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            value = ("" if cell.value is None else str(cell.value))
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

def export_excel(path):
    rows = load_rows()
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    # 01 — BARCHA MA'LUMOTLAR
    ws = wb.create_sheet("01_Barcha_Malumotlar")
    headers = ["id",
        "telegram_id",
        "name",
        "gender",
        "created_at",
        "age_group",
        "course",
        "residence",
        "marital_status",
        "ideal_marriage_age",
        "desired_children"]
    for i in range(1, 26):
        headers.append(f"q{i}")
    for i in range(1, 26):
        headers.append(f"q{i}_label")
    for i in range(1, 6):
        headers.append(f"top5_rank_{i}")
    for i in range(1, 7):
        headers.append(f"sd{i}")
    headers += ["sd_score", "sd_level", "main_factor", "custom_factor"]
    ws.append(headers)
    for row in rows:
        d = row["data"]
        values = [row["id"],
            row["telegram_id"],
            row["name"],
            row["gender"],
            row["created_at"],
            d.get("age_group", ""),
            d.get("course", ""),
            d.get("residence", ""),
            d.get("marital_status", ""),
            d.get("ideal_marriage_age", ""),
            d.get("desired_children", "")]
        answers = d.get("likert_answers", {})
        for key, label in CRITERIA:
            values.append(answers.get(key, ""))
        for key, label in CRITERIA:
            values.append(label)
        top5 = d.get("top5_ranks", [])
        for i in range(5):
            if i < len(top5):
                values.append(top5[i])
            else:
                values.append("")
        sd_answers = d.get("sd_answers", {})
        for i in range(1, 7):
            values.append(sd_answers.get(f"sd{i}", ""))
        values += [d.get("sd_score", ""), d.get("sd_level", ""), d.get("main_factor", ""), d.get(
            "custom_factor", "")]
        ws.append(values)
    style_sheet(ws)

    # 02 — DEMOGRAFIYA
    ws = wb.create_sheet("02_Demografiya")
    ws.append(["id", "gender", "age_group", "course", "residence", "marital_status",
        "ideal_marriage_age", "desired_children"])
    for row in rows:
        d = row["data"]
        ws.append([row["id"], row["gender"], d.get("age_group", ""), d.get("course", ""), d.get(
            "residence", ""), d.get("marital_status", ""), d.get("ideal_marriage_age", ""), d.get(
            "desired_children", "")])
    style_sheet(ws)

    # 03 — LIKERT TAHLILI
    ws = wb.create_sheet("03_Likert_Tahlil")
    ws.append(["q", "construct", "mezon", "N", "O‘rtacha", "Standart og‘ish", "0 soni", "1 soni",
        "2 soni", "3 soni", "3 ulushi (%)"])
    for idx, (key, label) in enumerate(CRITERIA, start=1):
        vals = []
        for row in rows:
            value = row["data"].get("likert_answers", {}).get(key)
            if isinstance(value, int):
                vals.append(value)
        n = len(vals)
        if n:
            mean = (sum(vals) / n)
        else:
            mean = ""
        if n > 1:
            variance = sum((x - mean) ** 2 for x in vals) / (n - 1)
            std = variance ** 0.5
        else:
            std = ""
        counts = [vals.count(i) for i in range(4)]
        if n:
            share3 = (counts[3] / n * 100)
        else:
            share3 = ""
        ws.append([f"q{idx}", key, label, n, round(mean, 4) if mean != "" else "", round(std, 4) if std != "" else "",
            *counts, round(share3, 2) if share3 != "" else ""])
    style_sheet(ws)

    # 04 — TOP5
    ws = wb.create_sheet("04_TOP5")
    ws.append(["rank", "mezon", "konstrukt", "N", "Ulush (%)"])
    total = len(rows)
    for rank in range(1, 6):
        counts = {}
        for row in rows:
            ranks = row["data"].get("top5_ranks", [])
            if len(ranks) >= rank:
                label = ranks[rank - 1]
                counts[label] = (counts.get(label, 0) + 1)
        for label, count in sorted(counts.items(), key=lambda x: (-x[1], x[0])):
            construct = next((k for k, l in CRITERIA if l == label), "")
            ws.append([rank, label, construct, count, round(count / total * 100, 2) if total else ""])
    style_sheet(ws)

    # 05 — SD
    ws = wb.create_sheet("05_SD")
    ws.append(["id", "gender", "sd1", "sd2", "sd3", "sd4", "sd5", "sd6", "sd_score", "sd_level"])
    for row in rows:
        d = row["data"]
        s = d.get("sd_answers", {})
        ws.append([row["id"], row["gender"], s.get("sd1", ""), s.get("sd2", ""), s.get("sd3", ""),
            s.get("sd4", ""), s.get("sd5", ""), s.get("sd6", ""), d.get("sd_score", ""), d.get(
            "sd_level", "")])
    style_sheet(ws)

    # 06 — GENDER TAHLILI
    ws = wb.create_sheet("06_Gender_Tahlil")
    ws.append(["q", "konstrukt", "mezon", "Erkak N", "Erkak o‘rtacha", "Ayol N", "Ayol o‘rtacha",
        "Farq (Erkak-Ayol)"])
    for idx, (key, label) in enumerate(CRITERIA, start=1):
        groups = {"Erkak": [], "Ayol": []}
        for row in rows:
            value = row["data"].get("likert_answers", {}).get(key)
            gender = row["gender"]
            if (isinstance(value, int) and gender in groups):
                groups[gender].append(value)
        means = {}
        for gender, vals in groups.items():
            if vals:
                means[gender] = sum(vals) / len(vals)
            else:
                means[gender] = ""
        if (means["Erkak"] != "" and means["Ayol"] != ""):
            diff = round(means["Erkak"] - means["Ayol"], 4)
        else:
            diff = ""
        ws.append([f"q{idx}", key, label, len(groups["Erkak"]), round(means["Erkak"], 4) if means[
            "Erkak"] != "" else "", len(groups["Ayol"]), round(means["Ayol"], 4) if means["Ayol"] != "" else "",
            diff])
    style_sheet(ws)

    # 07 — FAKTORLAR
    ws = wb.create_sheet("07_Faktorlar")
    ws.append(["Faktor", "N", "Ulush (%)"])
    factor_counts = {}
    for row in rows:
        factor = row["data"].get("main_factor", "")
        if factor:
            factor_counts[factor] = (factor_counts.get(factor, 0) + 1)
    for factor, count in sorted(factor_counts.items(), key=lambda x: -x[1]):
        ws.append([factor, count, round(count / total * 100, 2) if total else ""])
    style_sheet(ws)

    # 08 — SD TAHLILI
    ws = wb.create_sheet("08_SD_Tahlil")
    ws.append(["SD darajasi", "N", "Ulush (%)"])
    levels = {}
    for row in rows:
        level = row["data"].get("sd_level", "")
        if level:
            levels[level] = (levels.get(level, 0) + 1)
    for level in ["Past", "O‘rtacha", "Yuqori"]:
        count = levels.get(level, 0)
        ws.append([level, count, round(count / total * 100, 2) if total else ""])
    style_sheet(ws)

    # 09 — STATA
    ws = wb.create_sheet("09_Stata_Kodlar")
    stata_lines = ["* 25 ta Likert savoli: q1-q25",
        "* 0=Umuman muhim emas",
        "* 1=Unchalik muhim emas",
        "* 2=Muhim",
        "* 3=Juda muhim",
        "",
        "summarize q1-q25, detail",
        "alpha q1-q25",
        "pwcorr q1-q25, sig",
        "tabulate gender",
        "tabstat q1-q25, by(gender) stat(n mean sd)",
        "regress q1 q2 q3 q4 q5 q6 q7 q8 q9 q10",
        "regress q1 q2 q3 q4 q5 q6 q7 q8 q9 q10 q11 q12 q13 q14 q15 q16 q17 q18 q19 q20 q21 q22 q23 q24 q25"]
    for line in stata_lines:
        ws.append([line])
    style_sheet(ws)

    # 10 — CODEBOOK
    ws = wb.create_sheet("10_Codebook")
    ws.append(["O‘zgaruvchi", "Mazmuni", "Tip"])
    fixed = [("id", "Respondent identifikatori", "numeric"),
        ("telegram_id", "Telegram identifikatori", "numeric"),
        ("gender", "Jins", "categorical"),
        ("age_group", "Yosh guruhi", "categorical"),
        ("course", "Kurs", "categorical"),
        ("residence", "Doimiy yashash joyi", "categorical"),
        ("marital_status", "Oilaviy holat", "categorical"),
        ("ideal_marriage_age", "Nikohning maqbul yoshi", "numeric"),
        ("desired_children", "Istalgan farzandlar soni", "categorical"),
        ("sd_score", "Ijtimoiy maqbullik yig‘indi bali", "numeric"),
        ("sd_level", "Ijtimoiy maqbullik darajasi", "categorical"),
        ("main_factor", "Yakuniy asosiy omil", "categorical")]
    for item in fixed:
        ws.append(list(item))
    for i, (key, label) in enumerate(CRITERIA, start=1):
        ws.append([f"q{i}", label, "Likert 0–3"])
    for i, (key, label, expected) in enumerate(SD_QUESTIONS, start=1):
        ws.append([f"sd{i}", label, "binary"])
    for i in range(1, 6):
        ws.append([f"top5_rank_{i}", f"TOP-5 dagi {i}-o‘rin", "text"])
    style_sheet(ws)

    # 11 — UMUMIY STATISTIKA
    ws = wb.create_sheet("11_Umumiy_Statistika")
    total = len(rows)
    male = sum(1 for r in rows if r["gender"] == "Erkak")
    female = sum(1 for r in rows if r["gender"] == "Ayol")
    ws.append(["Ko‘rsatkich", "Qiymat"])
    ws.append(["Jami respondentlar", total])
    ws.append(["Erkak respondentlar", male])
    ws.append(["Ayol respondentlar", female])
    ws.append(["Test savollari soni", 25])
    ws.append(["TOP-5 mezonlar soni", 5])
    ws.append(["SD savollari soni", 6])
    style_sheet(ws)

    # 12 — XOM MA'LUMOT
    ws = wb.create_sheet("12_Xom_SQL_Malumot")
    ws.append(["id", "telegram_id", "name", "gender", "created_at", "data"])
    for row in rows:
        ws.append([row["id"], row["telegram_id"], row["name"], row["gender"], row["created_at"],
            repr(row["data"])])
    style_sheet(ws)
    wb.save(path)

# 12. ADMIN STATISTIKA

def admin_stats_text():
    rows = load_rows()
    total = len(rows)
    male = sum(1 for r in rows if r["gender"] == "Erkak")
    female = sum(1 for r in rows if r["gender"] == "Ayol")
    return ("ADMIN STATISTIKA\n\n" f"Jami respondent: {total}\n" f"Erkak: {male}\n"
        f"Ayol: {female}\n" "Test savollari: 25\n" "TOP-5: 5 ta tanlov\n" "SD: 6 ta savol")

# 13. MAIN

async def main():
    get_conn().close()

    # RENDER HEALTH CHECK
    app = web.Application()

    async def health(request):
        return web.Response(text="Bot ishlamoqda")
    app.router.add_get("/", health)
    runner = web.AppRunner(app)
    await runner.setup()
    port = int(os.environ.get("PORT", "8080"))
    site = web.TCPSite(runner, "0.0.0.0", port)
    await site.start()

    # TELEGRAM BOT
    bot = Bot(TOKEN)
    dp = Dispatcher()

    # START

    @dp.message(CommandStart())
    async def start(message: Message, state: FSMContext):
        if (message.from_user.id != ADMIN_ID and has_completed_response(message.from_user.id)):
            await state.clear()
            await message.answer("Siz ushbu so‘rovnomada " "avval qatnashgansiz.\n\n"
                "Ishtirokingiz uchun rahmat!")
            return
        await start_survey(message, state)

    # ADMIN

    @dp.message(Command("admin"))
    async def admin_panel(message: Message):
        if (message.from_user.id != ADMIN_ID):
            return
        await message.answer("ADMIN PANEL\n\n" "Kerakli amalni tanlang:", reply_markup=admin_keyboard(
            ))

    @dp.callback_query(F.data == "admin:stats")
    async def admin_stats(call: CallbackQuery):
        if (call.from_user.id != ADMIN_ID):
            await call.answer("Ruxsat yo‘q", show_alert=True)
            return
        await call.answer()
        await call.message.answer(admin_stats_text())

    @dp.callback_query(F.data == "admin:export")
    async def admin_export(call: CallbackQuery):
        if (call.from_user.id != ADMIN_ID):
            await call.answer("Ruxsat yo‘q", show_alert=True)
            return
        await call.answer("Excel tayyorlanmoqda...")
        try:
            export_excel(EXPORT_PATH)
            await call.message.answer_document(document=FSInputFile(EXPORT_PATH), caption=(
                "Excel tayyor.\n" "25 ta test, TOP-5, SD, " "gender tahlili, Stata "
                "va Codebook kiritilgan."))
        except Exception as e:
            await call.message.answer("Excel eksportida xatolik:\n" f"{type(e).__name__}: {e}")

    @dp.callback_query(F.data == "admin:reset")
    async def admin_reset(call: CallbackQuery):
        if (call.from_user.id != ADMIN_ID):
            await call.answer("Ruxsat yo‘q", show_alert=True)
            return
        delete_user_response(call.from_user.id)
        await call.answer("Test yozuvingiz o‘chirildi.")
        await call.message.answer("Mening testim o‘chirildi. "
            "Endi qayta boshlashingiz mumkin.")

    @dp.callback_query(F.data == "admin:reset_all")
    async def admin_reset_all(call: CallbackQuery):
        if (call.from_user.id != ADMIN_ID):
            await call.answer("Ruxsat yo‘q", show_alert=True)
            return
        delete_all_responses()
        await call.answer("Barcha ma’lumot o‘chirildi.", show_alert=True)
        await call.message.answer("Barcha respondent ma’lumotlari o‘chirildi.")

    @dp.callback_query(F.data == "admin:start")
    async def admin_start(call: CallbackQuery, state: FSMContext):
        if (call.from_user.id != ADMIN_ID):
            await call.answer("Ruxsat yo‘q", show_alert=True)
            return
        await call.answer()
        await start_survey(call.message, state)

    # ISM

    @dp.callback_query(Survey.name, F.data == "skip_name:0")
    async def skip_name(call: CallbackQuery, state: FSMContext):
        await state.update_data(name="")
        await state.set_state(Survey.gender)
        await call.answer()
        await call.message.edit_reply_markup(reply_markup=None)
        await call.message.answer("Jinsingizni tanlang:", reply_markup=gender_buttons())

    @dp.message(Survey.name)
    async def get_name(message: Message, state: FSMContext):
        name = (message.text or "").strip()
        if not name:
            await message.answer("Ismingizni yozing yoki "
                "«O‘tkazib yuborish» tugmasini bosing.")
            return
        await state.update_data(name=name)
        await state.set_state(Survey.gender)
        await message.answer("Jinsingizni tanlang:", reply_markup=gender_buttons())

    # JINS

    @dp.callback_query(Survey.gender, F.data.startswith("gender:"))
    async def get_gender(call: CallbackQuery, state: FSMContext):
        gender = call.data.split(":", 1)[1]
        await state.update_data(gender=gender, general_index=0, general_answers={},
            likert_answers={}, top5_selected=[], top5_ranks=[], sd_answers={}, sd_score=0)
        await call.answer()
        await call.message.edit_reply_markup(reply_markup=None)
        await send_general(call.message, state, 0)

    # DEMOGRAFIYA

    @dp.callback_query(Survey.general, F.data.startswith("general:"))
    async def get_general(call: CallbackQuery, state: FSMContext):
        data = await state.get_data()
        question_index = data.get("general_index", 0)
        option_index = int(call.data.split(":")[1])
        key, question, options = GENERAL_QUESTIONS[question_index]
        selected_answer = options[option_index]
        general_answers = data.get("general_answers", {})
        general_answers[key] = selected_answer
        next_index = question_index + 1
        await state.update_data(general_answers=general_answers, general_index=next_index)
        await call.answer()
        await call.message.edit_reply_markup(reply_markup=None)
        await send_general(call.message, state, next_index)

    # NIKOH YOSHI

    @dp.message(Survey.marriage_age)
    async def get_marriage_age(message: Message, state: FSMContext):
        value = (message.text or "").strip()
        try:
            age = int(value)
            if (age < 15 or age > 80):
                raise ValueError
        except ValueError:
            await message.answer("Iltimos, 15–80 oralig‘idagi " "yoshni raqam bilan kiriting.\n"
                "Masalan: 25")
            return
        await state.update_data(ideal_marriage_age=age)
        await send_likert(message, state, 0)

    # 25 TA TEST

    @dp.callback_query(Survey.likert, F.data.startswith("likert:"))
    async def get_likert(call: CallbackQuery, state: FSMContext):
        data = await state.get_data()
        index = data.get("likert_index", 0)
        value = int(call.data.split(":")[1])
        key, label = CRITERIA[index]
        answers = data.get("likert_answers", {})
        answers[key] = value
        await state.update_data(likert_answers=answers, likert_index=index + 1)
        await call.answer()
        await call.message.edit_reply_markup(reply_markup=None)
        await send_likert(call.message, state, index + 1)

    # TOP-5

    @dp.callback_query(Survey.top5, F.data.startswith("top5:"))
    async def get_top5(call: CallbackQuery, state: FSMContext):
        action = call.data.split(":", 1)[1]
        data = await state.get_data()
        selected = list(data.get("top5_selected", []))
        if action == "count":
            await call.answer(f"Tanlangan: " f"{len(selected)}/5", show_alert=True)
            return
        if action == "confirm":
            if len(selected) != 5:
                await call.answer("Avval 5 ta mezon tanlang.", show_alert=True)
                return
            ranks = [CRITERIA[i][1] for i in selected]
            await state.update_data(top5_ranks=ranks)
            await call.answer()
            await call.message.edit_reply_markup(reply_markup=None)
            await send_sd(call.message, state, 0)
            return
        try:
            idx = int(action)
        except ValueError:
            await call.answer()
            return
        if idx in selected:
            selected.remove(idx)
        else:
            if len(selected) >= 5:
                await call.answer("Faqat 5 ta mezon " "tanlash mumkin.", show_alert=True)
                return
            selected.append(idx)
        await state.update_data(top5_selected=selected)
        await call.answer()
        await call.message.edit_reply_markup(reply_markup=top5_keyboard(selected))

    # SD

    @dp.callback_query(Survey.sd_scale, F.data.startswith("sd:"))
    async def get_sd(call: CallbackQuery, state: FSMContext):
        data = await state.get_data()
        index = data.get("sd_index", 0)
        answer = int(call.data.split(":")[1])
        key, question, expected = SD_QUESTIONS[index]
        sd_answers = data.get("sd_answers", {})
        sd_answers[key] = answer
        score = data.get("sd_score", 0)
        if (bool(answer) == expected):
            score += 1
        await state.update_data(sd_answers=sd_answers, sd_score=score, sd_index=index + 1)
        await call.answer()
        await call.message.edit_reply_markup(reply_markup=None)
        await send_sd(call.message, state, index + 1)

    # YAKUNIY OMIL

    @dp.callback_query(Survey.main_factor, F.data.startswith("factor:"))
    async def get_factor(call: CallbackQuery, state: FSMContext):
        idx = int(call.data.split(":")[1])
        factor = MAIN_FACTORS[idx]
        await call.answer()
        await call.message.edit_reply_markup(reply_markup=None)
        if factor == "Boshqa":
            await state.set_state(Survey.custom_factor)
            await call.message.answer("Iltimos, siz uchun " "boshqa eng muhim " "omilni yozing:")
            return
        await state.update_data(main_factor=factor, custom_factor="")
        await finish_survey(call.message, state)

    @dp.message(Survey.custom_factor)
    async def get_custom_factor(message: Message, state: FSMContext):
        custom = (message.text or "").strip()
        if not custom:
            await message.answer("Iltimos, omilni yozib yuboring.")
            return
        await state.update_data(main_factor="Boshqa", custom_factor=custom)
        await finish_survey(message, state)

    # YAKUNLASH

    async def finish_survey(message, state):
        data = await state.get_data()
        sd_score = data.get("sd_score", 0)
        final_data = {"age_group": data.get("general_answers", {}).get("age_group", ""),
            "course": data.get("general_answers", {}).get("course", ""), "residence": data.get(
            "general_answers", {}).get("residence", ""), "marital_status": data.get(
            "general_answers", {}).get("marital_status", ""), "desired_children": data.get(
            "general_answers", {}).get("desired_children", ""), "ideal_marriage_age": data.get(
            "ideal_marriage_age", ""), "likert_answers": data.get("likert_answers", {}),
            "top5_ranks": data.get("top5_ranks", []), "sd_answers": data.get("sd_answers", {}),
            "sd_score": sd_score, "sd_level": calculate_sd_level(sd_score), "main_factor": data.get(
            "main_factor", ""), "custom_factor": data.get("custom_factor", "")}
        save_response(message.from_user.id, data.get("name", ""), data.get("gender", ""),
            final_data)
        await state.clear()
        await message.answer("So‘rovnoma yakunlandi!\n\n" "Ishtirokingiz uchun katta rahmat.\n"
            "Sizning javoblaringiz " "ilmiy-statistik tahlil uchun " "saqlab qo‘yildi.")

    # POLLING
    try:
        await dp.start_polling(bot)
    finally:
        await bot.session.close()
        await runner.cleanup()

# START
if __name__ == "__main__":
    import asyncio
    asyncio.run(main())
